#!/usr/bin/env python3
"""Build the 2026-10-16 V1.48 status handover by narrowly editing document.xml
and correcting the date in the existing footer.

The source DOCX is read-only. All package members other than word/document.xml
and word/footer2.xml are copied with their extracted payloads unchanged and
verified by SHA-256.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from lxml import etree
from docx import Document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML = "http://www.w3.org/XML/1998/namespace"
NS = {"w": W}
Q = lambda name: f"{{{W}}}{name}"
XML_SPACE = f"{{{XML}}}space"
EXPECTED_INPUT_SHA256 = "23a42e5f4cca40ef733edc9365c15bd0d5fdff647f8821fddec9c649e52b1455"
TARGET_PART = "word/document.xml"
FOOTER_PART = "word/footer2.xml"
TABLE_INDEX = 2

PARAGRAPH_REPLACEMENTS = [
    {
        "id": "subtitle",
        "index": 1,
        "old": "V1.48 A12合同对齐修订版  2026年10月9日",
        "new": "V1.48 工程进度与开发交接修订 2026年10月16日",
        "kind": "whole",
    },
    {
        "id": "opening_historical_scope",
        "index": 8,
        "old": "本次仅修订方案正文与验收门，不实施新功能；A01–A36及其历史PASS/FAIL/STOP/NOT_ESTABLISHED、历史数值与ORACLE_ACTION名称/冻结判决均不变。PG/VF/A12完成度及0个物理合格完整数据包记录不变。A17仅指job/checkpoint恢复；两份日期快照不证明当前状态，也不构成成员恢复证据。",
        "new": "V1.48（2026-10-09）仅修订方案正文与验收门，不实施新功能；A01–A36及其历史PASS/FAIL/STOP/NOT_ESTABLISHED、历史数值与ORACLE_ACTION名称/冻结判决均保持原样。本文记录的PG/VF/A12状态及0个物理合格完整数据包是截至2026-10-09的快照，不代表当前进度；当前状态见前置“2026年10月16日工程状态与交接”。A17仅指job/checkpoint恢复；V1.47/V1.48日期快照不证明当前状态，也不构成成员恢复证据。",
        "kind": "whole",
    },
    {
        "id": "document_version_and_cross_reference",
        "index": 92,
        "old": "本文件包含总纲、四岗位完整执行正文、软硬件、联调、验收、完整合成实验室、历史附录以及两个独立三维基础设施子项目方案，现行版本统一V1.47。总纲第2章给出最新统一任务、应用价值边界和评测合同；A6维护长期科研创新；A8把科研假设转成实验卡；A12与附录G负责合成、退化、预训练和诊断；附录I/J分别规定VesselForge与PG-Frag。附录A至F及历史实验段落继续作为冻结证据保留，旧数值与判决不能按V1.47的新目标直接改写。",
        "new": "本文件包含总纲、四岗位执行正文、软硬件、联调、验收、合成实验室、历史附录及两个独立三维基础设施子项目方案，现行版本统一为V1.48。合同对齐基线为2026-10-09；本次2026-10-16修订只补入工程状态、开发交接和资格边界，不改研究主线与历史证据。总纲第2章给出统一任务、应用价值边界和评测合同；A6维护长期科研创新；A8把研究假设转成实验卡；A12与附录G负责合成、退化、预训练和诊断；附录I/J分别规定VesselForge与PG-Frag。当前进度以本章前置的“2026年10月16日工程状态与交接”和1.2事实台账为准；带日期的旧状态保留为历史记录，A01–A36、MC01–MC07、历史数值及判决不得按新进度改写。",
        "kind": "whole",
    },
    {
        "id": "pause_scope",
        "index": 104,
        "old": "本轮范围仅为V1.47文档收口、研究候选转移与可行性审阅，不启动新开发、运行或新增数值结论。新列为P0的研究候选只表示后续研究优先级，不覆盖当前PG-Frag、VesselForge与A12的物理实现、资格和资源优先级，也不自动重开已停止工作；重新进入开发须由用户明确发起，并另行冻结范围与证据门槛。",
        "new": "V1.47文档收口时的历史范围记录：当时仅做文档收口、研究候选转移与可行性审阅，不启动新开发、运行或新增数值结论；P0仅表示研究优先级，不覆盖子项目物理实现、资格和资源优先级，也不自动重开已停止工作。\n截至2026-10-16，当前暂停范围允许归档GitHub代码、配置与测试，修订文档和总方案，并做离线交付检查；暂停新功能开发、求解器运行和新实验。v6零步preflight虽曾获授权，但暂停前未执行，暂停后不得启动；h、h/2、h/4 coupon campaign此前从未获准启动，两者授权状态不可混同。",
        "kind": "whole",
    },
    {
        "id": "a12_19_historical_label",
        "index": 510,
        "old": "A12.19 当前状态（2026-10-07）",
        "new": "A12.19 历史状态快照（截至2026-10-07；当前状态见前置交接）",
        "kind": "fragment",
    },
    {
        "id": "a12_19_percentages_historical",
        "index": 510,
        "old": "负责人给出的工程进度估计仍为 PG-Frag 25–35%、VesselForge 30–40%、A12 30–40%；这些比例是工程判断，不代表物理准确度或资格覆盖率。",
        "new": "负责人于2026-10-07给出的工程进度估计为 PG-Frag 25–35%、VesselForge 30–40%、A12 30–40%，仅保留为历史工程判断，不代表物理准确度或资格覆盖率；截至2026-10-16不发布新的综合百分比，当前进度按交付物、状态和资格门记录。",
        "kind": "fragment",
    },
    {
        "id": "a12_20_historical_label",
        "index": 513,
        "old": "A12.20 当前实施与资格边界（2026-10-07）",
        "new": "A12.20 历史实施与资格边界（截至2026-10-07；当前状态见前置交接）",
        "kind": "fragment",
    },
    {
        "id": "g31_historical_label",
        "index": 2183,
        "old": "G31.1 V1.47 收口核对（2026-10-07）",
        "new": "G31.1 V1.47 历史收口核对（截至2026-10-07；当前状态见前置交接）",
        "kind": "fragment",
    },
    {
        "id": "vesselforge_historical_label",
        "index": 2229,
        "old": "VesselForge 当前状态（2026-10-07）",
        "new": "VesselForge 历史状态快照（截至2026-10-07；当前状态见前置交接）",
        "kind": "fragment",
    },
    {
        "id": "pgfrag_historical_label",
        "index": 2356,
        "old": "PG-Frag 当前状态（2026-10-07）",
        "new": "PG-Frag 历史状态快照（截至2026-10-07；当前状态见前置交接）",
        "kind": "fragment",
    },
]

FRONT_SNAPSHOT = [
    ("Heading1", "2026年10月16日工程状态与交接", "heading"),
    ("Heading2", "当前交付与资格边界", "heading"),
    ("Normal", "本节记录截至2026年10月16日可核对的工程交付和边界。总体目标仍开放：真实日常与考古器型、多参数几何、来源感知材质、具有内容的连续纹样及外部资源精修仍需完成。软件检查、几何诊断和外观代理均不替代实物、物理或真实迁移资格。", "body"),
    ("Normal", "VesselForge已有Aberdeen v6、Met 18.56.44 v2来源指导候选及三个独立器型变体，并有STEP、表面、剖面、谱系和结构报告。v6 run05通过几何/来源身份诊断：86,624个正体积T4、CAD体积差0.240308%，44,182个外表三角保留16个source FID/role；它不是历史原器重建或物理资格。Apsida衍生扫描仍有7个pinched vertex links，仅可作视觉参考。Hasami Tokusa和素Jun有来源指导图案方向；Hacilar实际区域和motif尚未表达。Met 50899龙纹是烧成照片的局部外观代理，未复原初始颜料或完整纹样。所有外部资源仍待精修审核。", "body"),
    ("Normal", "A12的source_resolved_stationary_v1是可选的蓝白来源细节路径。单独的A12 source-resolution切片有55项聚焦检查通过；其geometry array与v6一致，但section hash不同，尚未接入已验收的native lineage链。该切片不是历史纹样恢复、校准釉料/烧成化学、完整整器连续纹样或真实迁移收益；照片裁切仍未完成精修。", "body"),
    ("Normal", "PG-Frag的旧native弱缝控制产生过2个闭合碎片和6个A12 RGB-D研究视图，但physical_pass=false、pretrain_eligible=false。v6 run05仍是网格/分区几何诊断。mother_binding工具检查通过不等于实际run05母器绑定验收；新母器自由stress-driven fracture尚未完成。完整合格物理包仍为0。", "body"),
    ("Heading2", "软件检查与未完成的物理门", "heading"),
    ("Normal", "最新聚焦软件检查为：VF 7文件66项PASS；PG 6文件55项PASS（含9项mother_binding）；A12 9文件套件90项tests collected、exit 0。三套共211项聚焦检查，不是全仓库测试，也不表示solver运行或物理资格。A12 source-resolution的55项切片是独立记录，不与PG的55项混算。", "body"),
    ("Normal", "v5全尺寸诊断网格曾完成零步初始化和单次100 ns零载荷成本控制，46个native步均无运动、损伤或能量变化；这不是断裂实验。v6 run05零步preflight曾获授权，但暂停前未执行。reference_v6_preflight当前只有输入、运行脚本和失败准备副本，没有native run receipt/result；暂停后不得启动该preflight。", "body"),
    ("Normal", "旧96-T4 coupon在step 652出现的负耗散可由所记录的负向native向量功解释，事件伴随未解析的方向反转和完全开口跃迁。检查到的Akantu v5.0.7源码中，getKineticEnergy可能调用assembleLumpedMatrix(\"M\")并重建质量矩阵；运行轮子为5.0.7-post1，源码一致性未确认，因此实际getter副作用尚未实测。C/D插入阶段getter前后mass与运动学快照仍是待做方案，尚未运行。", "body"),
    ("Heading2", "暂停范围与历史状态", "heading"),
    ("Normal", "当前暂停授权包括GitHub代码、配置和测试归档，文档与总方案修订，以及离线交付检查；禁止新功能开发、solver运行和新实验。h、h/2、h/4 coupon campaign从未获准启动，不能与曾获授权但未执行的v6零步preflight混为一谈。", "body"),
    ("Normal", "PG-Frag 25–35%、VesselForge 30–40%、A12 30–40%是2026-10-07的历史工程估计；本次不生成新的综合百分比。A01–A36及其历史PASS/FAIL/STOP/NOT_ESTABLISHED、数值和ORACLE_ACTION冻结判决保持原样；MC01–MC07继续为SPECIFIED / NOT_RUN。研究目标保持开放，完整合格物理包为0。", "body"),
]

A12_CURRENT = "2026-10-16状态补记：source_resolved_stationary_v1为opt-in来源细节路径；A12 source-resolution 55项专项检查与A12九文件套件90项测试是两组独立的软件记录。蓝白面板的技术组合不代表历史整器纹样恢复、初始颜料或烧成化学校准、完整纹样资产或真实检索迁移。其geometry array虽与v6一致，section hash不同，尚无native chain绑定。VF、PG、A12最新聚焦检查分别见前置交接快照；它们不构成物理资格。"
VF_CURRENT = "2026-10-16状态补记：Aberdeen v6与Met 18.56.44 v2仍为来源指导候选，另有三个独立器型变体及STEP、表面、剖面、谱系和结构报告。v6 run05通过几何/来源身份诊断（86,624个正体积T4、0.240308% CAD体积差、16个source FID/role），不构成历史原器重建或物理资格。日常与考古器型覆盖、参数与材料校准、真实语义纹样及外部资源精修仍需逐项完成。Apsida衍生扫描仅限视觉参考，Hacilar motif尚未表达；详见前置当前状态。"
PG_CURRENT = "2026-10-16状态补记：旧native弱缝控制的2个闭合碎片和6个A12 RGB-D视图仍属研究控制，physical_pass=false、pretrain_eligible=false。PG 6文件55项聚焦检查通过，其中mother_binding占9项；该软件验证不等于实际run05母器绑定验收。新母器stress-driven fracture未完成。v6零步preflight曾获授权但暂停前未执行，暂停后不得启动；h、h/2、h/4 coupon campaign从未获准启动。getter副作用仅由v5.0.7源码审阅推测，运行轮子5.0.7-post1的一致性与实际效果未验证；详见EXTRINSIC_EVENT_MATH.md。"

TABLE2_ROWS = [
    [
        "VesselForge",
        "Aberdeen v6、Met 18.56.44 v2来源候选及三个器型变体；run05为86,624 T4、0.240308%体积差、16个source FID/role；VF 7文件66项检查PASS。",
        "几何/来源身份诊断；非历史原器重建或物理资格。多参数器型、纹样、材料与外部资源精修仍未签收。",
    ],
    [
        "PG-Frag",
        "旧弱缝控制有2个闭合碎片和6个RGB-D研究视图；PG 6文件55项检查PASS，含9项mother_binding；v6零步preflight未执行。",
        "旧控制physical_pass=false；新母器stress-driven fracture和实际run05母器绑定未验收。v6网格仅为几何诊断。",
    ],
    [
        "A12 / SyntheticLab",
        "source_resolved_stationary_v1为opt-in；source-resolution切片55项聚焦检查通过；A12九文件套件90项tests collected、exit 0。",
        "技术外观路径；非完整整器纹样、历史颜料/化学复原或真实迁移。资源精修与物理资格仍未完成。",
    ],
]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def direct_body_paragraphs(body: etree._Element) -> list[etree._Element]:
    return [child for child in body if child.tag == Q("p")]


def direct_body_tables(body: etree._Element) -> list[etree._Element]:
    return [child for child in body if child.tag == Q("tbl")]


def paragraph_text(p: etree._Element) -> str:
    return "".join(node.text or "" for node in p.xpath(".//w:t", namespaces=NS))


def set_xml_space(node: etree._Element, text: str) -> None:
    if text and (text[0].isspace() or text[-1].isspace()):
        node.set(XML_SPACE, "preserve")


def replace_span(p: etree._Element, old: str, new: str, label: str) -> None:
    nodes = p.xpath(".//w:t", namespaces=NS)
    whole = "".join(n.text or "" for n in nodes)
    count = whole.count(old)
    if count != 1:
        raise ValueError(f"{label}: expected one exact anchor, found {count}: {old[:80]!r}")
    start = whole.find(old)
    end = start + len(old)
    spans: list[tuple[int, int, etree._Element]] = []
    cursor = 0
    for node in nodes:
        value = node.text or ""
        spans.append((cursor, cursor + len(value), node))
        cursor += len(value)
    active = [(a, b, n) for a, b, n in spans if b > start and a < end]
    if not active:
        raise ValueError(f"{label}: anchor mapped to no text node")
    first_a, first_b, first = active[0]
    last_a, last_b, last = active[-1]
    if first is last:
        prior = (first.text or "")[:start - first_a]
        post = (first.text or "")[end - first_a:]
        first.text = prior + new + post
        set_xml_space(first, first.text or "")
    else:
        first.text = (first.text or "")[:start - first_a] + new
        set_xml_space(first, first.text or "")
        for _, _, node in active[1:-1]:
            node.text = ""
        last.text = (last.text or "")[end - last_a:]
        set_xml_space(last, last.text or "")


def xml_para_text_donor(donor: etree._Element, text: str, style_id: str | None = None) -> etree._Element:
    p = etree.Element(Q("p"))
    ppr = donor.find(Q("pPr"))
    if ppr is not None:
        p.append(copy.deepcopy(ppr))
    if style_id:
        new_ppr = p.find(Q("pPr"))
        if new_ppr is None:
            new_ppr = etree.Element(Q("pPr"))
            p.insert(0, new_ppr)
        pstyle = new_ppr.find(Q("pStyle"))
        if pstyle is None:
            pstyle = etree.Element(Q("pStyle"))
            new_ppr.insert(0, pstyle)
        pstyle.set(Q("val"), style_id)
    run = etree.SubElement(p, Q("r"))
    donor_rpr = donor.find("./" + Q("r") + "/" + Q("rPr"))
    if donor_rpr is not None:
        run.append(copy.deepcopy(donor_rpr))
    t = etree.SubElement(run, Q("t"))
    t.text = text
    set_xml_space(t, text)
    return p


def add_page_break_before(p: etree._Element) -> None:
    ppr = p.find(Q("pPr"))
    if ppr is None:
        ppr = etree.Element(Q("pPr"))
        p.insert(0, ppr)
    if ppr.find(Q("pageBreakBefore")) is None:
        el = etree.Element(Q("pageBreakBefore"))
        # CT_PPr places pageBreakBefore after keepNext/keepLines and before spacing.
        idx = 0
        for i, child in enumerate(ppr):
            if child.tag in {Q("keepNext"), Q("keepLines")}:
                idx = i + 1
            elif child.tag == Q("spacing"):
                break
        ppr.insert(idx, el)


def c14n(element: etree._Element) -> bytes:
    return etree.tostring(element, method="c14n", exclusive=False, with_comments=True)


def cell_text(cell: etree._Element) -> str:
    return "".join(node.text or "" for node in cell.xpath(".//w:t", namespaces=NS))


def cell_text_rows(table: etree._Element) -> list[list[str]]:
    result = []
    for row in table.findall(Q("tr")):
        result.append([cell_text(c) for c in row.findall(Q("tc"))])
    return result


def set_cell_text(cell: etree._Element, text: str) -> None:
    paragraphs = cell.findall(Q("p"))
    if not paragraphs:
        p = etree.SubElement(cell, Q("p"))
        r = etree.SubElement(p, Q("r"))
        t = etree.SubElement(r, Q("t"))
        t.text = text
        set_xml_space(t, text)
        return
    p = paragraphs[0]
    nodes = p.xpath(".//w:t", namespaces=NS)
    if not nodes:
        run = etree.SubElement(p, Q("r"))
        t = etree.SubElement(run, Q("t"))
        t.text = text
        set_xml_space(t, text)
    else:
        nodes[0].text = text
        set_xml_space(nodes[0], text)
        for node in nodes[1:]:
            node.text = ""
    for extra in paragraphs[1:]:
        cell.remove(extra)


def ensure_repeat_header(row: etree._Element) -> None:
    trpr = row.find(Q("trPr"))
    if trpr is None:
        trpr = etree.Element(Q("trPr"))
        row.insert(0, trpr)
    if trpr.find(Q("tblHeader")) is None:
        marker = etree.Element(Q("tblHeader"))
        marker.set(Q("val"), "true")
        after = -1
        for i, child in enumerate(trpr):
            if child.tag in {Q("cantSplit"), Q("trHeight")}:
                after = i
        trpr.insert(after + 1, marker)


def prepare_new_row(template: etree._Element, values: list[str]) -> etree._Element:
    row = copy.deepcopy(template)
    trpr = row.find(Q("trPr"))
    if trpr is None:
        trpr = etree.Element(Q("trPr"))
        row.insert(0, trpr)
    # Rows auto-expand; do not copy a fixed row height from the source template.
    for height in list(trpr.findall(Q("trHeight"))):
        trpr.remove(height)
    if trpr.find(Q("cantSplit")) is None:
        trpr.insert(0, etree.Element(Q("cantSplit")))
    cells = row.findall(Q("tc"))
    if len(cells) != len(values):
        raise ValueError(f"Table row has {len(cells)} cells but {len(values)} values were supplied")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def paragraph_insertions(original_paragraphs: list[etree._Element]) -> dict[int, list[etree._Element]]:
    donor_normal = original_paragraphs[92]
    donor_h1 = original_paragraphs[91]
    donor_h2 = original_paragraphs[99]
    donor_syn = original_paragraphs[510]
    insertions: dict[int, list[etree._Element]] = {}

    front = []
    for style_id, text, kind in FRONT_SNAPSHOT:
        donor = donor_h1 if style_id == "Heading1" else donor_h2 if style_id == "Heading2" else donor_normal
        para = xml_para_text_donor(donor, text, style_id)
        if text == "2026年10月16日工程状态与交接":
            add_page_break_before(para)
        front.append(para)
    insertions[9] = front

    insertions[514] = [xml_para_text_donor(donor_syn, A12_CURRENT)]
    insertions[2230] = [xml_para_text_donor(donor_syn, VF_CURRENT)]
    insertions[2357] = [xml_para_text_donor(donor_syn, PG_CURRENT)]
    return insertions


def expected_paragraphs(source_paragraph_texts: list[str], insertions: dict[int, list[etree._Element]]) -> list[str]:
    replacements_by_index: dict[int, list[dict[str, str]]] = {}
    for op in PARAGRAPH_REPLACEMENTS:
        replacements_by_index.setdefault(op["index"], []).append(op)
    expected: list[str] = []
    for i, source_text in enumerate(source_paragraph_texts):
        expected.extend(paragraph_text(x) for x in insertions.get(i, []))
        text = source_text
        for op in replacements_by_index.get(i, []):
            old = op["old"]
            new = op.get("new", "")
            if text.count(old) != 1:
                raise ValueError(f"Expected paragraph {i} anchor for {op['id']} exactly once")
            text = text.replace(old, new, 1)
        expected.append(text)
    expected.extend(paragraph_text(x) for x in insertions.get(len(source_paragraph_texts), []))
    return expected


def write_package(source: Path, target: Path, modified_parts: dict[str, bytes]) -> dict[str, Any]:
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(target.name + ".building")
    if temp.exists():
        temp.unlink()
    try:
        with zipfile.ZipFile(source, "r") as zin:
            with zipfile.ZipFile(temp, "w", allowZip64=True) as zout:
                zout.comment = zin.comment
                for info in zin.infolist():
                    payload = modified_parts.get(info.filename, zin.read(info))
                    zout.writestr(copy.copy(info), payload)
        os.replace(temp, target)
    finally:
        if temp.exists():
            temp.unlink()
    return {"target": str(target), "sha256": sha256_file(target), "bytes": target.stat().st_size}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt-dir", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    source = args.input.resolve()
    target = args.output.resolve()
    receipt_dir = args.receipt_dir.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if target.exists() and not args.overwrite:
        raise FileExistsError(f"Refusing to overwrite {target}; use --overwrite only for this task-owned output")
    source_sha = sha256_file(source)
    if source_sha != EXPECTED_INPUT_SHA256:
        raise ValueError(f"Source hash mismatch: {source_sha}")
    mirror = Path("relicweave-synlab/docs/master/20261009/RelicWeave_V1.48_ContractAligned.docx")
    mirror_sha = sha256_file(mirror) if mirror.is_file() else None
    if mirror_sha != EXPECTED_INPUT_SHA256:
        raise ValueError(f"20261009 archive mirror hash mismatch: {mirror_sha}")

    with zipfile.ZipFile(source, "r") as zin:
        source_parts = {i.filename: zin.read(i) for i in zin.infolist()}
        source_infos = list(zin.infolist())
        if TARGET_PART not in source_parts:
            raise ValueError(f"Missing {TARGET_PART}")
        if FOOTER_PART not in source_parts:
            raise ValueError(f"Missing {FOOTER_PART}")

    # Update only the printed date text in the footer. A raw, unique UTF-8
    # replacement preserves every field run and all other footer markup byte-for-byte.
    old_footer_date = "2026年10月9日".encode("utf-8")
    new_footer_date = "2026年10月16日".encode("utf-8")
    source_footer = source_parts[FOOTER_PART]
    if source_footer.count(old_footer_date) != 1:
        raise ValueError("Expected exactly one 2026年10月9日 date in word/footer2.xml")
    modified_footer_xml = source_footer.replace(old_footer_date, new_footer_date, 1)
    if modified_footer_xml.count(new_footer_date) != 1 or modified_footer_xml.count(old_footer_date):
        raise ValueError("Footer date replacement verification failed")
    if modified_footer_xml.replace(new_footer_date, old_footer_date, 1) != source_footer:
        raise ValueError("Footer payload changed beyond the date text")

    # Capture the existing body's payload fingerprint before replacing the output.
    # The rebuild must produce exactly the same document.xml as root's rendered v1.
    previous_output_document_xml_sha256 = None
    if target.is_file():
        with zipfile.ZipFile(target, "r") as previous_zip:
            previous_output_document_xml_sha256 = sha256_bytes(previous_zip.read(TARGET_PART))
    parser_xml = etree.XMLParser(remove_blank_text=False, resolve_entities=False, huge_tree=True)
    root = etree.fromstring(source_parts[TARGET_PART], parser_xml)
    body = root.find(Q("body"))
    if body is None:
        raise ValueError("word/document.xml has no w:body")
    original_paragraphs = direct_body_paragraphs(body)
    original_tables = direct_body_tables(body)
    if len(original_paragraphs) != 2479 or len(original_tables) != 195:
        raise ValueError(f"Unexpected source structure: {len(original_paragraphs)} paragraphs, {len(original_tables)} tables")
    source_para_texts = [paragraph_text(p) for p in original_paragraphs]
    table2 = original_tables[TABLE_INDEX]
    table2_original_rows = cell_text_rows(table2)
    if len(table2_original_rows) != 8:
        raise ValueError(f"Unexpected 1.2 table row count: {len(table2_original_rows)}")

    # Replace only anchored text spans in explicitly approved paragraphs.
    for op in PARAGRAPH_REPLACEMENTS:
        p = original_paragraphs[op["index"]]
        replace_span(p, op["old"], op["new"], op["id"])

    # Add the current status section before the existing V1.47 navigation paragraph.
    insertions = paragraph_insertions(original_paragraphs)
    for original_index, new_paras in sorted(insertions.items(), reverse=True):
        target_p = original_paragraphs[original_index] if original_index < len(original_paragraphs) else None
        if target_p is None:
            for new_p in new_paras:
                body.insert(len(body) - 1, new_p)
        else:
            pos = body.index(target_p)
            for offset, new_p in enumerate(new_paras):
                body.insert(pos + offset, new_p)

    # Append compact project rows to 1.2; preserve all pre-existing row text.
    table2 = direct_body_tables(body)[TABLE_INDEX]
    table_rows = table2.findall(Q("tr"))
    if len(table_rows) != 8:
        raise ValueError("1.2 table changed unexpectedly before row append")
    ensure_repeat_header(table_rows[0])
    template = table_rows[-1]
    for values in TABLE2_ROWS:
        table2.append(prepare_new_row(template, values))

    modified_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # Verify planned paragraph insertion and text edits before creating the ZIP.
    expected_paragraph_texts = expected_paragraphs(source_para_texts, insertions)
    output_paragraphs_xml = direct_body_paragraphs(body)
    output_paragraph_texts = [paragraph_text(p) for p in output_paragraphs_xml]
    if output_paragraph_texts != expected_paragraph_texts:
        for i, (a, b) in enumerate(zip(output_paragraph_texts, expected_paragraph_texts)):
            if a != b:
                raise ValueError(f"Output paragraph sequence mismatch at {i}: {a[:120]!r} != {b[:120]!r}")
        raise ValueError(f"Paragraph count mismatch: {len(output_paragraph_texts)} != {len(expected_paragraph_texts)}")

    # Verify old table cells were retained, other tables are unchanged, and the new table repeats its header.
    final_tables = direct_body_tables(body)
    if len(final_tables) != len(original_tables):
        raise ValueError("Table count changed")
    final_table2 = final_tables[TABLE_INDEX]
    final_rows = cell_text_rows(final_table2)
    if final_rows[:len(table2_original_rows)] != table2_original_rows:
        raise ValueError("Pre-existing 1.2 table row text changed")
    if len(final_rows) != len(table2_original_rows) + 3:
        raise ValueError("Expected exactly 3 appended 1.2 rows")
    header_rows = final_table2.findall(Q("tr"))
    header_has_repeat = header_rows[0].find(f"./{Q('trPr')}/{Q('tblHeader')}") is not None
    if not header_has_repeat:
        raise ValueError("1.2 table header is not configured to repeat")
    if any(row.find(f"./{Q('trPr')}/{Q('trHeight')}") is not None for row in header_rows[-3:]):
        raise ValueError("New status rows must auto-expand; fixed row height found")
    unchanged_tables = []
    changed_tables = []
    for i, (before, after) in enumerate(zip(original_tables, final_tables)):
        if i == TABLE_INDEX:
            continue
        if c14n(before) != c14n(after):
            changed_tables.append(i)
        else:
            unchanged_tables.append(i)
    if changed_tables:
        raise ValueError(f"Unexpected table changes outside T{TABLE_INDEX:03}: {changed_tables}")

    # Verify important history/contract anchors remain present after edits.
    final_body_text = "\n".join(output_paragraph_texts)
    table_text = "\n".join(" | ".join(row) for row in final_rows)
    required_phrases = [
        "A01–A36", "MC01–MC07", "NOT_ESTABLISHED", "ORACLE_ACTION",
        "25–35%", "30–40%", "截至2026-10-07", "physical_pass=false",
        "SPECIFIED / NOT_RUN", "0个物理合格完整数据包",
    ]
    missing = [phrase for phrase in required_phrases if phrase not in final_body_text + "\n" + table_text]
    if missing:
        raise ValueError(f"Required history/status phrases missing: {missing}")

    # Build the output package. Only document.xml and the corrected footer2.xml may differ.
    if previous_output_document_xml_sha256 is not None and previous_output_document_xml_sha256 != sha256_bytes(modified_xml):
        raise ValueError("Generated document.xml differs from the current rendered-v1 body payload")
    modified_parts = {TARGET_PART: modified_xml, FOOTER_PART: modified_footer_xml}
    package_result = write_package(source, target, modified_parts)
    if sha256_file(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("Source file changed during build")
    with zipfile.ZipFile(target, "r") as zout:
        output_infos = list(zout.infolist())
        output_parts = {i.filename: zout.read(i) for i in output_infos}
    if [i.filename for i in source_infos] != [i.filename for i in output_infos]:
        raise ValueError("ZIP member order or names changed")
    if set(source_parts) != set(output_parts):
        raise ValueError("ZIP member set changed")
    part_records = []
    changed_parts = []
    for name in source_parts:
        before = source_parts[name]
        after = output_parts[name]
        same = before == after
        if not same:
            changed_parts.append(name)
        part_records.append({
            "name": name,
            "source_bytes": len(before),
            "output_bytes": len(after),
            "source_sha256": sha256_bytes(before),
            "output_sha256": sha256_bytes(after),
            "payload_unchanged": same,
        })
    expected_changed_parts = [name for name in source_parts if name in {TARGET_PART, FOOTER_PART}]
    if changed_parts != expected_changed_parts:
        raise ValueError(f"Expected only {expected_changed_parts} payloads to change; got {changed_parts}")
    if output_parts[TARGET_PART] != modified_xml:
        raise ValueError("Output document.xml differs from the planned body payload")
    if output_parts[FOOTER_PART] != modified_footer_xml:
        raise ValueError("Output footer2.xml differs from the planned date-only payload")
    for part_name in source_parts:
        if part_name not in {TARGET_PART, FOOTER_PART} and source_parts[part_name] != output_parts[part_name]:
            raise ValueError(f"Unapproved package part changed: {part_name}")
    footer_text = output_parts[FOOTER_PART].decode("utf-8")
    if "2026年10月16日" not in footer_text or "2026年10月9日" in footer_text:
        raise ValueError("Final footer date text is not updated")
    for field_instruction in (b" PAGE ", b" NUMPAGES "):
        if source_footer.count(field_instruction) != output_parts[FOOTER_PART].count(field_instruction):
            raise ValueError(f"Footer field instruction changed: {field_instruction!r}")
    for field_type in (b"begin", b"separate", b"end"):
        marker = b'w:fldCharType="' + field_type + b'"'
        if source_footer.count(marker) != output_parts[FOOTER_PART].count(marker):
            raise ValueError(f"Footer field boundary count changed: {field_type.decode()}")
    footer_field_run_counts = {
        "PAGE_instruction": output_parts[FOOTER_PART].count(b" PAGE "),
        "NUMPAGES_instruction": output_parts[FOOTER_PART].count(b" NUMPAGES "),
        "begin_markers": output_parts[FOOTER_PART].count(b'w:fldCharType="begin"'),
        "separate_markers": output_parts[FOOTER_PART].count(b'w:fldCharType="separate"'),
        "end_markers": output_parts[FOOTER_PART].count(b'w:fldCharType="end"'),
    }

    # Re-open with python-docx to confirm the produced package parses and has the expected structure.
    parsed = Document(str(target))
    if len(parsed.paragraphs) != len(output_paragraph_texts) or len(parsed.tables) != 195:
        raise ValueError("python-docx output readback structure mismatch")
    if parsed.paragraphs[1].text != "V1.48 工程进度与开发交接修订 2026年10月16日":
        raise ValueError("Subtitle readback mismatch")
    if parsed.paragraphs[9].text != "2026年10月16日工程状态与交接":
        raise ValueError("Front snapshot insertion/readback position mismatch")
    if len(parsed.tables[TABLE_INDEX].rows) != 11:
        raise ValueError("1.2 table readback row count mismatch")

    receipt_dir.mkdir(parents=True, exist_ok=True)
    other_table_fingerprint = hashlib.sha256(b"".join(c14n(final_tables[i]) for i in unchanged_tables)).hexdigest()
    history_receipt = {
        "source_document": str(source),
        "source_sha256": source_sha,
        "output_document": str(target),
        "output_sha256": package_result["sha256"],
        "source_body_paragraphs": len(original_paragraphs),
        "output_body_paragraphs": len(output_paragraph_texts),
        "paragraph_sequence_matches_expected": True,
        "paragraph_replacements_applied": [op["id"] for op in PARAGRAPH_REPLACEMENTS],
        "inserted_paragraph_count": len(output_paragraph_texts) - len(original_paragraphs),
        "historical_and_contract_markers_preserved": {phrase: phrase in final_body_text + "\n" + table_text for phrase in required_phrases},
        "table_count_source_output": [len(original_tables), len(final_tables)],
        "table_1_2_index": TABLE_INDEX,
        "table_1_2_rows_source_output": [len(table2_original_rows), len(final_rows)],
        "table_1_2_original_rows_text_unchanged": True,
        "table_1_2_new_rows": TABLE2_ROWS,
        "table_1_2_repeating_header": True,
        "table_1_2_new_rows_auto_expand_no_fixed_height": True,
        "other_original_tables_unchanged_count": len(unchanged_tables),
        "other_original_tables_changed_indices": changed_tables,
        "other_original_tables_c14n_fingerprint_sha256": other_table_fingerprint,
        "current_output_document_xml_sha256_before_rebuild": previous_output_document_xml_sha256,
        "output_document_xml_byte_identical_to_rendered_v1": previous_output_document_xml_sha256 == sha256_bytes(modified_xml),
        "footer_date_updated": {
            "part": FOOTER_PART,
            "old": "2026年10月9日",
            "new": "2026年10月16日",
            "replacement_scope": "date text bytes only; remaining footer bytes preserved",
            "page_and_numpages_field_runs_preserved": True,
        },
        "render_status": "PENDING_ROOT_FULL_RENDER_AND_VISUAL_QA",
        "delivery_status": "NOT_FINAL_UNTIL_ROOT_RENDER_QA_PASSES",
    }
    part_receipt = {
        "source_document": str(source),
        "source_sha256": source_sha,
        "mirror_20261009_sha256": mirror_sha,
        "output_document": str(target),
        "output_sha256": package_result["sha256"],
        "output_bytes": package_result["bytes"],
        "zip_member_count": len(source_parts),
        "zip_member_names_and_order_unchanged": True,
        "declared_part_scope": [TARGET_PART, FOOTER_PART],
        "all_parts_outside_declared_scope_byte_identical": True,
        "document_xml_byte_identical_to_current_render_v1": previous_output_document_xml_sha256 == sha256_bytes(modified_xml),
        "document_xml_payload_sha256": sha256_bytes(output_parts[TARGET_PART]),
        "footer_date_update": {
            "part": FOOTER_PART,
            "source_date": "2026年10月9日",
            "output_date": "2026年10月16日",
            "source_payload_sha256": sha256_bytes(source_footer),
            "output_payload_sha256": sha256_bytes(output_parts[FOOTER_PART]),
            "only_date_bytes_changed": output_parts[FOOTER_PART].replace(new_footer_date, old_footer_date, 1) == source_footer,
            "page_and_numpages_field_runs_preserved": True,
            "field_instruction_and_boundary_counts": footer_field_run_counts,
        },
        "changed_parts": changed_parts,
        "parts": part_records,
        "source_still_matches_expected_hash": sha256_file(source) == EXPECTED_INPUT_SHA256,
        "settings_updateFields_modified": False,
        "render_status": "PENDING_ROOT_FULL_RENDER_AND_VISUAL_QA",
    }
    (receipt_dir / "anchor_history_verification.json").write_text(json.dumps(history_receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (receipt_dir / "source_part_preservation.json").write_text(json.dumps(part_receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    command = (
        f'& "{Path(sys.executable).resolve()}" -X utf8 "{Path(__file__).resolve()}" '
        f'--input "{source}" --output "{target}" --receipt-dir "{receipt_dir}" --overwrite'
    )
    readme = f'''# RelicWeave V1.48 工程进度修订构建记录

本目录中的 `{target.name}` 是从只读 V1.48 归档副本构建的独立修订件。源 SHA-256：`{source_sha}`。输出 SHA-256：`{package_result["sha256"]}`。

修订涉及 `word/document.xml` 与 `word/footer2.xml`：前者增加截至 2026-10-16 的工程状态与交接快照，更新 1.2 当前事实台账，补充 A12、VesselForge、PG-Frag 的当前状态交叉引用，并把 2026-10-07/09 的状态明确保留为历史；后者只把页脚日期改为 2026年10月16日，PAGE/NUMPAGES 字段运行保持不变。没有更改研究主线、A01–A36 历史判决、MC01–MC07 合同状态或任何物理实验结果。V1.48 版号保留；原 2026-10-09 文档未覆盖。

`source_part_preservation.json` 逐项记录 ZIP part payload 的 SHA-256；仅 `word/document.xml` 和 `word/footer2.xml` 属于声明修订范围，其余 part 的解压内容均与源件字节一致。`anchor_history_verification.json` 记录原始段落锚点、历史文本、1.2 表原行与其余 194 张表的未改验证，并确认正文 XML 与渲染 v1 完全相同、页脚字段运行保留。

构建命令（从工作区根目录运行）：

```powershell
& "{Path(sys.executable).resolve()}" -X utf8 "{Path(__file__).resolve()}" --input "{source}" --output "{target}" --receipt-dir "{receipt_dir}" --overwrite
```

状态：**等待主代理完成全量渲染、逐页 PNG 检查和最终布局 QA**。当前不将本文件报告为已通过视觉验收或最终发布。
'''
    (receipt_dir / "README.md").write_text(readme, encoding="utf-8")

    run_record = {
        "command": command,
        "source_sha256": source_sha,
        "output": package_result,
        "declared_part_scope": [TARGET_PART, FOOTER_PART],
        "changed_parts": changed_parts,
        "document_xml_byte_identical_to_current_render_v1": previous_output_document_xml_sha256 == sha256_bytes(modified_xml),
        "footer_date_updated": "2026年10月9日 -> 2026年10月16日",
        "page_and_numpages_field_runs_preserved": True,
        "paragraph_anchor_count": len(PARAGRAPH_REPLACEMENTS),
        "inserted_paragraph_count": len(output_paragraph_texts) - len(original_paragraphs),
        "table2_rows_appended": 3,
        "render_status": "PENDING_ROOT_FULL_RENDER_AND_VISUAL_QA",
    }
    log_path = Path(__file__).resolve().parent / "build_run.json"
    log_path.write_text(json.dumps(run_record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(run_record, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
