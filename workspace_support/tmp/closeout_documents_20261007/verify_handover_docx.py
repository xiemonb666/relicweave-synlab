from pathlib import Path
import re, hashlib, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT=Path(r'C:\Users\xiemo\Desktop\ShardForge')
md_path=ROOT/'docs/handover/20261007/HANDOVER.md'
docx_path=ROOT/'docs/handover/20261007/RelicWeave_A12_VesselForge_PGFrag_Handover_20261007.docx'
expected_md_sha='f226a153f8aed2da04e07298f1277cd74071370d52cf74a4967702c2792f7a53'
actual_md_sha=hashlib.sha256(md_path.read_bytes()).hexdigest()
if actual_md_sha!=expected_md_sha: raise SystemExit(f'MD changed after build: {actual_md_sha}')
lines=md_path.read_text(encoding='utf-8-sig').splitlines()
expected=[]; expected_code=[]; i=0; in_code=False; code=[]
def unmark(s):
    s=re.sub(r'\*\*(.*?)\*\*',r'\1',s)
    s=re.sub(r'`(.*?)`',r'\1',s)
    s=re.sub(r'\*([^*]+)\*',r'\1',s)
    return s
while i<len(lines):
    line=lines[i]
    if line.startswith('~~~'):
        if not in_code: in_code=True; code=[]
        else: expected_code.append('\n'.join(code)); expected.append('\n'.join(code)); in_code=False; code=[]
        i+=1; continue
    if in_code: code.append(line); i+=1; continue
    if line.startswith('|') and line.rstrip().endswith('|'):
        while i<len(lines) and lines[i].startswith('|') and lines[i].rstrip().endswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if vals and not all(re.fullmatch(r':?-{3,}:?',x or '-') for x in vals):
                expected.extend(unmark(v) for v in vals)
            i+=1
        continue
    if not line.strip(): i+=1; continue
    if line.startswith('# '): expected.append(unmark(line[2:].strip()))
    elif line.startswith('## '): expected.append(unmark(line[3:].strip()))
    elif line.startswith('### '): expected.append(unmark(line[4:].strip()))
    elif line.startswith('#### '): expected.append(unmark(line[5:].strip()))
    elif line.startswith('- '): expected.append(unmark(line[2:].strip()))
    else: expected.append(unmark(line.strip()))
    i+=1
if in_code: raise SystemExit('Markdown has an open code fence')
doc=Document(docx_path)
actual=[]; code_actual=[]
for child in doc._element.body.iterchildren():
    if child.tag==qn('w:p'):
        p=Paragraph(child,doc._body)
        if p.text: actual.append(p.text)
        ppr=child.pPr
        shd=ppr.find(qn('w:shd')) if ppr is not None else None
        if shd is not None and shd.get(qn('w:fill'))=='F6F7F9': code_actual.append(p.text)
    elif child.tag==qn('w:tbl'):
        tbl=Table(child,doc._body)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text: actual.append(p.text)
def compact(parts): return ''.join(c for c in ''.join(parts) if not c.isspace())
expected_flat=compact(expected); actual_flat=compact(actual)
if expected_flat!=actual_flat:
    limit=min(len(expected_flat),len(actual_flat)); i=next((j for j in range(limit) if expected_flat[j]!=actual_flat[j]),limit)
    raise SystemExit(f'Visible text mismatch at compact char {i}: expected_len={len(expected_flat)} actual_len={len(actual_flat)} expected={expected_flat[max(0,i-30):i+60]!r} actual={actual_flat[max(0,i-30):i+60]!r}')
if expected_code!=code_actual:
    raise SystemExit(f'Code block mismatch: expected {len(expected_code)}, actual {len(code_actual)}')
hashes=set(re.findall(r'(?<![0-9a-fA-F])[0-9a-fA-F]{64}(?![0-9a-fA-F])',md_path.read_text(encoding='utf-8-sig')))
missing=sorted(h for h in hashes if h not in actual_flat)
if missing: raise SystemExit(f'Hashes absent from DOCX visible text: {missing}')
sec=doc.sections[0]
if round(sec.page_width.inches,2)!=8.5 or round(sec.page_height.inches,2)!=11.0: raise SystemExit('Page size is not US Letter 8.5x11')
header=doc.sections[0].header.paragraphs[0]
header_colors=[r.font.color.rgb for r in header.runs if r.text]
if any(c is not None and str(c).upper()!='000000' for c in header_colors): raise SystemExit(f'Header is not black: {header_colors}')
print('markdown_sha256',actual_md_sha)
print('docx_sha256',hashlib.sha256(docx_path.read_bytes()).hexdigest(),'bytes',docx_path.stat().st_size)
print('compact_visible_characters',len(actual_flat),'markdown_full_hashes',len(hashes),'docx_matched_hashes',len(hashes))
print('code_blocks',len(expected_code),'code_lines',sum(x.count('\n')+1 for x in expected_code))
print('page_size_inches',round(sec.page_width.inches,2),round(sec.page_height.inches,2),'header_colors',[str(c) for c in header_colors])
