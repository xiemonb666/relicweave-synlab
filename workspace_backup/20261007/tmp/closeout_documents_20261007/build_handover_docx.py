from __future__ import annotations
from pathlib import Path
import argparse, hashlib, re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(r'C:\Users\xiemo\Desktop\ShardForge')
DEFAULT_MD=ROOT/'docs/handover/20261007/HANDOVER.md'
DEFAULT_OUT=ROOT/'docs/handover/20261007/RelicWeave_A12_VesselForge_PGFrag_Handover_20261007.docx'
FONT='Microsoft YaHei'
CODE_FONT='Consolas'
NAVY='203864'
LIGHT='F3F6FA'
GRID='D9D9D9'
BLACK='000000'

def sha(p:Path)->str: return hashlib.sha256(p.read_bytes()).hexdigest()
def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill); shd.set(qn('w:val'),'clear')
def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn(f'w:{name}'))
        if node is None: node=OxmlElement(f'w:{name}'); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def set_table_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),GRID)
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)
def add_field(run, field):
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=f' {field} '
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    text=OxmlElement('w:t'); text.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,instr,separate,text,end])
def set_run_fonts(run, size=9.5, bold=False, italic=False, color=BLACK, name=FONT):
    run.font.name=name; run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor.from_string(color)
    rPr=run._element.get_or_add_rPr(); rFonts=rPr.rFonts
    if rFonts is None: rFonts=OxmlElement('w:rFonts'); rPr.insert(0,rFonts)
    rFonts.set(qn('w:ascii'),name); rFonts.set(qn('w:hAnsi'),name); rFonts.set(qn('w:eastAsia'),FONT)
def add_styled_text(paragraph, text, base_size=9.5, bold=False, italic=False, color=BLACK, name=FONT):
    pieces=re.split(r'(?<![0-9A-Fa-f])([0-9A-Fa-f]{64})(?![0-9A-Fa-f])',text)
    for piece in pieces:
        if not piece: continue
        run=paragraph.add_run()
        set_run_fonts(run,base_size,bold=bold,italic=italic,color=color,name=name)
        if re.fullmatch(r'[0-9A-Fa-f]{64}',piece):
            run.add_text(piece[:32]); run.add_break(); run.add_text(piece[32:])
        else:
            run.add_text(piece)
def add_inline(paragraph, text, base_size=9.5, table=False):
    token=re.compile(r'(\*\*.*?\*\*|`.*?`|\*[^*]+\*)')
    pos=0
    for m in token.finditer(text):
        if m.start()>pos:
            add_styled_text(paragraph,text[pos:m.start()],base_size)
        item=m.group(0)
        if item.startswith('**'):
            add_styled_text(paragraph,item[2:-2],base_size,bold=True)
        elif item.startswith('`'):
            add_styled_text(paragraph,item[1:-1],base_size-0.3,name=CODE_FONT)
        else:
            add_styled_text(paragraph,item[1:-1],base_size,italic=True)
        pos=m.end()
    if pos<len(text):
        add_styled_text(paragraph,text[pos:],base_size)
def configure(doc):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.32); sec.footer_distance=Inches(.32)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name=FONT; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.08
    for name,size in [('Heading 1',16),('Heading 2',12.5),('Heading 3',10.8)]:
        st=styles[name]; st.font.name=FONT; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(BLACK)
        st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); st.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 9); st.paragraph_format.space_after=Pt(4); st.paragraph_format.keep_with_next=True
    title=styles['Title']; title.font.name=FONT; title.font.size=Pt(22); title.font.bold=True; title.font.color.rgb=RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); title.paragraph_format.space_after=Pt(5); title.paragraph_format.keep_with_next=True
    subtitle=styles['Subtitle']; subtitle.font.name=FONT; subtitle.font.size=Pt(10.5); subtitle.font.color.rgb=RGBColor.from_string('404040')
    subtitle._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); subtitle.paragraph_format.space_after=Pt(7)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=header.add_run('RELICWEAVE  |  V1.47 HANDOVER'); set_run_fonts(r,7.5,color=BLACK)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('工程交接  ·  2026-10-07  ·  '); set_run_fonts(r,7.5,color='555555')
    r=footer.add_run(); set_run_fonts(r,7.5,color='555555'); add_field(r,'PAGE')
    doc.core_properties.title='RelicWeave A12 VesselForge PG-Frag Handover 2026-10-07'
    doc.core_properties.subject='Implementation closeout, saved evidence, constraints, and next-step boundary'
    doc.core_properties.keywords='RelicWeave; A12; VesselForge; PG-Frag; V1.47'
def add_table(doc, rows):
    n=len(rows[0]); table=doc.add_table(rows=1,cols=n); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    if n==2: widths=[Inches(1.32),Inches(5.47)]
    elif n==3: widths=[Inches(1.12),Inches(2.83),Inches(2.84)]
    elif n==4: widths=[Inches(.82),Inches(2.1),Inches(1.75),Inches(2.12)]
    else: widths=[Inches(6.79/n)]*n
    for ci,w in enumerate(widths): table.columns[ci].width=w
    set_table_borders(table)
    for ri,rowvals in enumerate(rows):
        row=table.rows[0] if ri==0 else table.add_row()
        for ci,val in enumerate(rowvals):
            cell=row.cells[ci]; cell.width=widths[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
            if ri==0: set_cell_shading(cell,NAVY)
            elif ri%2==0: set_cell_shading(cell,LIGHT)
            para=cell.paragraphs[0]; para.paragraph_format.space_after=Pt(1); para.paragraph_format.line_spacing=1.0
            add_inline(para,val,8.2,table=True)
            for run in para.runs:
                if ri==0: run.font.color.rgb=RGBColor.from_string('FFFFFF'); run.font.bold=True
        if ri==0: set_repeat_table_header(row)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def build(md_path:Path,out_path:Path,expected_sha:str):
    actual=sha(md_path)
    if actual.lower()!=expected_sha.lower(): raise SystemExit(f'Markdown snapshot mismatch: expected {expected_sha}, got {actual}')
    lines=md_path.read_text(encoding='utf-8-sig').splitlines()
    doc=Document(); configure(doc)
    i=0; in_code=False; code_lines=[]
    while i<len(lines):
        line=lines[i]
        if line.startswith('~~~'):
            if not in_code:
                in_code=True; code_lines=[]
            else:
                p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2); p.paragraph_format.line_spacing=1.0
                r=p.add_run('\n'.join(code_lines)); set_run_fonts(r,8.0,name=CODE_FONT,color='202020')
                pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F6F7F9'); shd.set(qn('w:val'),'clear'); pPr.append(shd)
                in_code=False; code_lines=[]
            i+=1; continue
        if in_code: code_lines.append(line); i+=1; continue
        if line.startswith('|') and line.rstrip().endswith('|'):
            group=[]
            while i<len(lines) and lines[i].startswith('|') and lines[i].rstrip().endswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not vals or all(re.fullmatch(r':?-{3,}:?',x or '-') for x in vals): i+=1; continue
                group.append(vals); i+=1
            if group: add_table(doc,group)
            continue
        if not line.strip(): i+=1; continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title'); add_inline(p,line[2:].strip(),22); i+=1; continue
        if line.startswith('## '):
            p=doc.add_paragraph(style='Heading 1'); add_inline(p,line[3:].strip(),16); i+=1; continue
        if line.startswith('### '):
            p=doc.add_paragraph(style='Heading 2'); add_inline(p,line[4:].strip(),12.5); i+=1; continue
        if line.startswith('#### '):
            p=doc.add_paragraph(style='Heading 3'); add_inline(p,line[5:].strip(),10.8); i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.13); add_inline(p,line[2:].strip(),9.2); i+=1; continue
        p=doc.add_paragraph(); add_inline(p,line.strip(),9.5); i+=1
    if in_code: raise SystemExit('Unclosed Markdown code block')
    out_path.parent.mkdir(parents=True,exist_ok=True); doc.save(out_path)
    print(f'markdown_sha256={actual}')
    print(f'output={out_path}')
    print(f'output_bytes={out_path.stat().st_size}')
    print(f'output_sha256={sha(out_path)}')
if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--markdown',type=Path,default=DEFAULT_MD); ap.add_argument('--output',type=Path,default=DEFAULT_OUT); ap.add_argument('--expected-sha',required=True)
    a=ap.parse_args(); build(a.markdown,a.output,a.expected_sha)
